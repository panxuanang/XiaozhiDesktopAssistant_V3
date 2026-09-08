from __future__ import annotations

import json
import os
import sqlite3
import time
from pathlib import Path
from typing import Iterable

from .paths import file_index_db_path
from .tools.files import resolve_user_path

_TEXT_EXTS = {".txt", ".md", ".csv", ".json", ".log", ".ini", ".yaml", ".yml"}
_DOC_EXTS = {".docx", ".pdf", ".xlsx", ".xlsm", ".pptx"}


def _conn() -> sqlite3.Connection:
    conn = sqlite3.connect(file_index_db_path(), timeout=20)
    conn.row_factory = sqlite3.Row
    conn.execute(
        """CREATE TABLE IF NOT EXISTS docs(
            path TEXT PRIMARY KEY,
            name TEXT NOT NULL,
            ext TEXT NOT NULL,
            mtime REAL NOT NULL,
            size INTEGER NOT NULL,
            content TEXT NOT NULL,
            indexed_at REAL NOT NULL
        )"""
    )
    conn.execute("CREATE INDEX IF NOT EXISTS idx_docs_name ON docs(name)")
    return conn


def _extract(path: Path, max_chars: int = 60000) -> str:
    ext = path.suffix.lower()
    try:
        if ext in _TEXT_EXTS:
            return path.read_text(encoding="utf-8", errors="ignore")[:max_chars]
        if ext == ".docx":
            from docx import Document
            doc = Document(path)
            return "\n".join(p.text for p in doc.paragraphs)[:max_chars]
        if ext == ".pdf":
            from pypdf import PdfReader
            chunks = []
            for page in PdfReader(path).pages[:80]:
                chunks.append(page.extract_text() or "")
                if sum(map(len, chunks)) >= max_chars:
                    break
            return "\n".join(chunks)[:max_chars]
        if ext in {".xlsx", ".xlsm"}:
            import pandas as pd
            book = pd.ExcelFile(path)
            parts = []
            for sheet in book.sheet_names[:8]:
                df = pd.read_excel(path, sheet_name=sheet, nrows=300)
                parts.append(f"[工作表:{sheet}]\n" + df.to_csv(index=False))
                if sum(map(len, parts)) >= max_chars:
                    break
            return "\n".join(parts)[:max_chars]
        if ext == ".pptx":
            from pptx import Presentation
            prs = Presentation(path)
            lines = []
            for slide in prs.slides[:80]:
                for shape in slide.shapes:
                    if hasattr(shape, "text") and shape.text:
                        lines.append(shape.text)
                if sum(map(len, lines)) >= max_chars:
                    break
            return "\n".join(lines)[:max_chars]
    except Exception:
        return ""
    return ""


def index_roots(roots: str | Iterable[str], max_files: int = 12000) -> dict[str, int]:
    if isinstance(roots, str):
        root_items = [x.strip() for x in roots.replace("\n", ";").split(";") if x.strip()]
    else:
        root_items = list(roots)
    seen = indexed = skipped = 0
    with _conn() as conn:
        for root_text in root_items:
            root = resolve_user_path(root_text)
            if not root.exists():
                continue
            for current, dirs, files in os.walk(root):
                dirs[:] = [d for d in dirs if not d.startswith(".") and d.lower() not in {"node_modules", ".git", "$recycle.bin"}]
                for name in files:
                    seen += 1
                    if seen > max_files:
                        return {"seen": seen - 1, "indexed": indexed, "skipped": skipped}
                    path = Path(current) / name
                    ext = path.suffix.lower()
                    if ext not in _TEXT_EXTS | _DOC_EXTS:
                        skipped += 1
                        continue
                    try:
                        stat = path.stat()
                        row = conn.execute("SELECT mtime,size FROM docs WHERE path=?", (str(path),)).fetchone()
                        if row and float(row[0]) == stat.st_mtime and int(row[1]) == stat.st_size:
                            skipped += 1
                            continue
                        content = _extract(path)
                        conn.execute(
                            "INSERT INTO docs(path,name,ext,mtime,size,content,indexed_at) VALUES(?,?,?,?,?,?,?) "
                            "ON CONFLICT(path) DO UPDATE SET name=excluded.name,ext=excluded.ext,mtime=excluded.mtime,size=excluded.size,content=excluded.content,indexed_at=excluded.indexed_at",
                            (str(path), path.name, ext, stat.st_mtime, stat.st_size, content, time.time()),
                        )
                        indexed += 1
                    except (OSError, PermissionError):
                        skipped += 1
        conn.commit()
    return {"seen": seen, "indexed": indexed, "skipped": skipped}


def search_index(query: str, limit: int = 30) -> list[dict[str, object]]:
    terms = [t for t in query.strip().split() if t]
    if not terms:
        return []
    clauses = []
    args: list[object] = []
    for term in terms:
        clauses.append("(lower(name) LIKE ? OR lower(content) LIKE ?)")
        needle = f"%{term.lower()}%"
        args += [needle, needle]
    args.append(limit)
    sql = "SELECT path,name,ext,mtime,size,substr(content,1,800) AS preview FROM docs WHERE " + " AND ".join(clauses) + " ORDER BY mtime DESC LIMIT ?"
    with _conn() as conn:
        rows = conn.execute(sql, args).fetchall()
    return [dict(r) for r in rows]


def find_recent_files(roots: str, extensions: str = ".xlsx,.xls,.docx,.pdf,.pptx", within_hours: int = 72, limit: int = 30) -> list[dict[str, object]]:
    exts = {x.strip().lower() if x.strip().startswith(".") else "." + x.strip().lower() for x in extensions.replace("，", ",").split(",") if x.strip()}
    cutoff = time.time() - within_hours * 3600
    results: list[dict[str, object]] = []
    for root_text in [x.strip() for x in roots.split(";") if x.strip()]:
        root = resolve_user_path(root_text)
        if not root.exists():
            continue
        for current, dirs, files in os.walk(root):
            dirs[:] = [d for d in dirs if not d.startswith(".")]
            for name in files:
                p = Path(current) / name
                if p.suffix.lower() not in exts:
                    continue
                try:
                    st = p.stat()
                except OSError:
                    continue
                if st.st_mtime >= cutoff:
                    results.append({"path": str(p), "name": p.name, "modified": st.st_mtime, "size": st.st_size})
    results.sort(key=lambda x: float(x["modified"]), reverse=True)
    return results[:limit]
