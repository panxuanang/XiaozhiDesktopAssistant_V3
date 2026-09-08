from __future__ import annotations

import json
from pathlib import Path
from typing import Any

from pptx import Presentation
from pptx.util import Inches, Pt

from ..action_log import record
from .files import resolve_user_path


def create_presentation(path: str, title: str, slides_json: str | list[dict[str, Any]], subtitle: str = "") -> str:
    """Create a simple business PPT locally. slides_json is [{title, bullets:[...]}]."""
    target = resolve_user_path(path)
    if target.suffix.lower() != ".pptx":
        target = target.with_suffix(".pptx")
    target.parent.mkdir(parents=True, exist_ok=True)
    slides = json.loads(slides_json) if isinstance(slides_json, str) else slides_json
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)

    cover = prs.slides.add_slide(prs.slide_layouts[0])
    cover.shapes.title.text = title
    cover.placeholders[1].text = subtitle
    cover.shapes.title.text_frame.paragraphs[0].font.size = Pt(30)

    for item in slides[:30]:
        slide = prs.slides.add_slide(prs.slide_layouts[1])
        slide.shapes.title.text = str(item.get("title", ""))
        body = slide.placeholders[1].text_frame
        body.clear()
        bullets = item.get("bullets", [])
        if isinstance(bullets, str):
            bullets = [x.strip() for x in bullets.split("\n") if x.strip()]
        for i, bullet in enumerate(bullets[:12]):
            p = body.paragraphs[0] if i == 0 else body.add_paragraph()
            p.text = str(bullet)
            p.level = 0
            p.font.size = Pt(20)
    prs.save(target)
    record("create_presentation", {"path": str(target), "slides": len(slides) + 1})
    return str(target)


def create_ppt_from_text(path: str, source_text: str, title: str = "工作汇报", max_slides: int = 8) -> str:
    """Use the configured API only for slide outlining, then generate the PPT locally."""
    from ..llm import AIClient
    ai = AIClient()
    prompt = f"""把下面材料整理成不超过 {max_slides} 页的中文汇报PPT提纲。
只输出 JSON 数组，每项格式：{{"title":"页标题","bullets":["要点1","要点2"]}}。
不要输出Markdown。
材料：\n{source_text[:30000]}"""
    raw = ai.complete(prompt, system="你是办公汇报PPT结构设计助手，只输出严格JSON。", max_output_tokens=2500)
    raw = raw.strip().strip("`")
    if raw.startswith("json"):
        raw = raw[4:].strip()
    slides = json.loads(raw)
    return create_presentation(path, title, slides)


def read_presentation(path: str, max_chars: int = 30000) -> str:
    target = resolve_user_path(path)
    prs = Presentation(target)
    lines: list[str] = []
    for i, slide in enumerate(prs.slides, start=1):
        lines.append(f"[第{i}页]")
        for shape in slide.shapes:
            if hasattr(shape, "text") and shape.text:
                lines.append(shape.text)
        if sum(map(len, lines)) >= max_chars:
            break
    return "\n".join(lines)[:max_chars]


def create_ppt_from_excel_local(excel_path: str, output_path: str = "", title: str = "Excel 数据汇报") -> str:
    """Create a compact PPT from deterministic Excel statistics without any AI call."""
    from .office import excel_local_analysis
    src = resolve_user_path(excel_path)
    analysis = excel_local_analysis(str(src))
    lines = [x.strip() for x in analysis.splitlines() if x.strip()]
    overview = lines[:4]
    detail = lines[4:9]
    slides = [
        {"title": "数据概览", "bullets": overview or ["已读取本地 Excel 数据"]},
        {"title": "关键指标", "bullets": detail or overview},
        {"title": "下一步", "bullets": ["关注最低值和未达标项", "结合业务背景复盘异常原因", "确认后可继续生成更详细的AI分析"]},
    ]
    target = output_path or str(src.with_name(src.stem + "_汇报.pptx"))
    return create_presentation(target, title, slides, subtitle=src.name)
