from __future__ import annotations

import json
from pathlib import Path
from typing import Any
from copy import copy

import pandas as pd
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt

from ..action_log import record
from ..config import load_settings
from .files import resolve_user_path


def _apply_default_word_style(doc: Document) -> None:
    style = doc.styles["Normal"]
    style.font.name = "宋体"
    style.font.size = Pt(12)


def create_word_document(path: str, title: str, content: str, author: str = "") -> str:
    target = resolve_user_path(path)
    if target.suffix.lower() != ".docx":
        target = target.with_suffix(".docx")
    target.parent.mkdir(parents=True, exist_ok=True)
    doc = Document()
    _apply_default_word_style(doc)
    heading = doc.add_paragraph()
    heading.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = heading.add_run(title.strip())
    run.bold = True
    run.font.name = "黑体"
    run.font.size = Pt(20)
    for block in [b.strip() for b in content.replace("\r\n", "\n").split("\n")]:
        if not block:
            doc.add_paragraph()
            continue
        p = doc.add_paragraph(block)
        p.paragraph_format.first_line_indent = Pt(24)
        p.paragraph_format.line_spacing = 1.5
    if author:
        p = doc.add_paragraph(author)
        p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    doc.save(target)
    record("create_word_document", {"path": str(target), "title": title}, detail=f"chars={len(content)}")
    return str(target)


def write_material(requirement: str, path: str = "桌面/材料.docx", title: str = "") -> str:
    from ..llm import AIClient
    ai = AIClient()
    prompt = f"""请根据以下要求直接撰写一份可交付的中文材料。\n\n用户要求：{requirement}\n\n要求：结构完整、事实不编造、语言符合用户要求；不要解释写作过程。"""
    content = ai.complete(prompt, system="你是一名专业中文材料写作助手，擅长通知、总结、汇报、方案、会议纪要和商务材料。")
    if not title:
        title = ai.complete(f"给下面材料拟一个简洁正式的标题，只输出标题：\n{content[:4000]}", max_output_tokens=80).strip("《》 \n")
    saved = create_word_document(path, title, content)
    return f"材料已生成并保存：{saved}"


def read_word_document(path: str, max_chars: int = 30000) -> str:
    target = resolve_user_path(path)
    doc = Document(target)
    text = "\n".join(p.text for p in doc.paragraphs)
    return text[:max_chars]


def rewrite_word_document(path: str, instruction: str, output_path: str = "") -> str:
    from ..llm import AIClient
    src = resolve_user_path(path)
    original = read_word_document(str(src))
    ai = AIClient()
    revised = ai.complete(
        f"修改要求：{instruction}\n\n原文：\n{original}",
        system="你是专业文稿编辑。保持原文核心事实，不添加未经提供的事实，直接输出修改后的完整正文。",
    )
    target = resolve_user_path(output_path) if output_path else src.with_name(src.stem + "_修改版.docx")
    saved = create_word_document(str(target), src.stem, revised)
    return f"已生成修改版：{saved}"


def excel_profile(path: str, max_rows: int = 50000) -> dict[str, Any]:
    target = resolve_user_path(path)
    book = pd.ExcelFile(target)
    result: dict[str, Any] = {"file": str(target), "sheets": {}}
    for sheet in book.sheet_names[:20]:
        df = pd.read_excel(target, sheet_name=sheet, nrows=max_rows)
        info: dict[str, Any] = {
            "rows": int(len(df)),
            "columns": [str(c) for c in df.columns],
            "missing": {str(k): int(v) for k, v in df.isna().sum().items() if int(v) > 0},
            "sample": df.head(8).where(pd.notna(df.head(8)), None).to_dict(orient="records"),
        }
        num = df.select_dtypes(include="number")
        if not num.empty:
            desc = num.describe().round(4).to_dict()
            info["numeric_summary"] = desc
            info["numeric_totals"] = {str(k): (None if pd.isna(v) else float(v)) for k, v in num.sum(numeric_only=True).items()}
        cat_summary: dict[str, Any] = {}
        for col in df.select_dtypes(exclude="number").columns[:12]:
            values = df[col].astype(str).value_counts(dropna=True).head(10)
            cat_summary[str(col)] = {str(k): int(v) for k, v in values.items()}
        if cat_summary:
            info["top_values"] = cat_summary
        result["sheets"][sheet] = info
    return result


def analyze_excel(path: str, question: str = "请分析这份表格的关键结论、异常和建议", report_path: str = "") -> str:
    from ..llm import AIClient
    from ..llm.client import json_for_ai
    profile = excel_profile(path)
    settings = load_settings()
    compact = json_for_ai(profile, max_chars=settings.safety.max_file_chars_for_ai)
    ai = AIClient()
    answer = ai.complete(
        f"用户问题：{question}\n\n下面是本地程序对 Excel 做的结构化统计摘要：\n{compact}",
        system="你是资深数据分析师。基于提供的统计结果作答；不要臆造原表中没有的数据；用清晰中文给出结论、证据和可执行建议。",
    )
    record("analyze_excel", {"path": path, "question": question}, detail=answer[:500])
    if report_path:
        saved = create_word_document(report_path, "Excel 数据分析报告", answer)
        return answer + f"\n\n分析报告已保存：{saved}"
    return answer


def excel_preview(path: str, sheet: str = "", rows: int = 20) -> str:
    target = resolve_user_path(path)
    df = pd.read_excel(target, sheet_name=sheet if sheet else 0, nrows=max(1, min(rows, 100)))
    return df.to_csv(index=False)



def _find_excel_header(ws, header: str, max_rows: int = 30):
    needle = str(header).strip()
    for row in ws.iter_rows(min_row=1, max_row=min(max_rows, ws.max_row)):
        for cell in row:
            if str(cell.value or "").strip() == needle:
                return cell.row, cell.column
    raise ValueError(f"没有找到 Excel 表头：{header}")


def excel_calculate_row_totals(
    path: str,
    source_headers: str = "服饰,数码,家居,食品,美妆",
    total_header: str = "总销售额（万元）",
    sheet: str = "",
) -> dict[str, Any]:
    """按指定列逐行求和并直接写回 Excel 数值，避免公式缓存导致后续分析读不到结果。"""
    from openpyxl import load_workbook

    target = resolve_user_path(path)
    wb = load_workbook(target)
    ws = wb[sheet] if sheet else wb[wb.sheetnames[0]]

    headers = [h.strip() for h in source_headers.replace("，", ",").split(",") if h.strip()]
    if not headers:
        raise ValueError("source_headers 不能为空")
    found = [_find_excel_header(ws, h) for h in headers]
    header_rows = {r for r, _ in found}
    if len(header_rows) != 1:
        raise ValueError("源列不在同一表头行")
    source_header_row = next(iter(header_rows))
    try:
        total_row, total_col = _find_excel_header(ws, total_header)
    except ValueError:
        total_row, total_col = source_header_row, ws.max_column + 1
        ws.cell(total_row, total_col).value = total_header
    header_rows.add(total_row)
    if len(header_rows) != 1:
        raise ValueError("源列和总计列不在同一表头行")
    header_row = total_row
    source_cols = [c for _, c in found]

    results: list[dict[str, Any]] = []
    for r in range(header_row + 1, ws.max_row + 1):
        # 第一列为空或进入分析区后停止；允许中间没有数据的行跳过
        label = ws.cell(r, 1).value
        numeric_values = []
        valid = True
        for c in source_cols:
            v = ws.cell(r, c).value
            if v in (None, ""):
                valid = False
                break
            try:
                numeric_values.append(float(v))
            except (TypeError, ValueError):
                valid = False
                break
        if not valid:
            if results and (label in (None, "") or str(label).strip().startswith("AI")):
                break
            continue
        total = sum(numeric_values)
        out_cell = ws.cell(r, total_col)
        out_cell.value = int(total) if float(total).is_integer() else round(total, 4)
        results.append({"row": r, "label": str(label or r), "total": out_cell.value})

    if not results:
        raise ValueError("没有找到可计算的数据行")
    wb.save(target)
    record("excel_calculate_row_totals", {"path": str(target), "sheet": ws.title, "rows": len(results), "total_header": total_header})
    return {"file": str(target), "sheet": ws.title, "rows_updated": len(results), "results": results}


def excel_write_analysis(
    path: str,
    sales_top3: str,
    below_target: str,
    category_performance: str,
    summary: str,
    sheet: str = "",
) -> str:
    """将四项简短分析写入模板中的 AI 分析区，保留原有样式。"""
    from openpyxl import load_workbook
    from openpyxl.styles import Alignment

    target = resolve_user_path(path)
    wb = load_workbook(target)
    ws = wb[sheet] if sheet else wb[wb.sheetnames[0]]
    mapping = {
        "销售前三": sales_top3,
        "低于目标门店": below_target,
        "品类表现": category_performance,
        "一句话总结": summary,
    }
    written = []
    for label, text in mapping.items():
        located = False
        for row in ws.iter_rows():
            for cell in row:
                if str(cell.value or "").strip() == label:
                    # 模板的标签区是 A:B 合并，正文区从 C 开始；通用情况下写到右侧第二列。
                    target_col = cell.column + 2
                    out = ws.cell(cell.row, target_col)
                    out.value = str(text)
                    try:
                        out.alignment = Alignment(horizontal=out.alignment.horizontal, vertical="center", text_rotation=out.alignment.text_rotation, wrap_text=True, shrink_to_fit=out.alignment.shrink_to_fit, indent=out.alignment.indent)
                    except Exception:
                        pass
                    written.append(f"{out.coordinate}={text}")
                    located = True
                    break
            if located:
                break
        if not located:
            raise ValueError(f"没有找到分析区标签：{label}")
    wb.save(target)
    record("excel_write_analysis", {"path": str(target), "sheet": ws.title}, detail=" | ".join(written)[:1000])
    return f"分析已写入：{target}（{ws.title}）"


def excel_local_analysis(path: str, sheet: str = "") -> str:
    """Produce a deterministic local-only Excel summary. No AI/API call."""
    target = resolve_user_path(path)
    df = pd.read_excel(target, sheet_name=sheet if sheet else 0)
    lines = [f"文件：{target.name}", f"数据行：{len(df)}，列数：{len(df.columns)}", "列：" + "、".join(map(str, df.columns))]
    num = df.select_dtypes(include="number")
    if not num.empty:
        totals = num.sum(numeric_only=True).sort_values(ascending=False)
        lines.append("数值列合计：" + "；".join(f"{k}={v:.2f}" for k, v in totals.head(8).items()))
        for col in num.columns[:8]:
            series = num[col].dropna()
            if len(series):
                idx_max = series.idxmax(); idx_min = series.idxmin()
                label_col = df.columns[0]
                lines.append(f"{col}：最高 {df.loc[idx_max, label_col]}={series.loc[idx_max]:.2f}；最低 {df.loc[idx_min, label_col]}={series.loc[idx_min]:.2f}；均值 {series.mean():.2f}")
    missing = df.isna().sum()
    missing = missing[missing > 0]
    if len(missing):
        lines.append("缺失值：" + "；".join(f"{k}={int(v)}" for k, v in missing.items()))
    return "\n".join(lines)


def _excel_headers(path: str, sheet: str = "") -> list[str]:
    target = resolve_user_path(path)
    df = pd.read_excel(target, sheet_name=sheet if sheet else 0, nrows=5)
    return [str(c) for c in df.columns]


def _guess_numeric_headers(path: str, sheet: str = "") -> list[str]:
    target = resolve_user_path(path)
    df = pd.read_excel(target, sheet_name=sheet if sheet else 0, nrows=30)
    return [str(c) for c in df.select_dtypes(include="number").columns]


def excel_plan_from_instruction(path: str, instruction: str, sheet: str = "", allow_ai: bool = True) -> dict[str, Any]:
    """Translate natural language into a small local spreadsheet plan.

    Common instructions are recognized with rules. Only ambiguous tasks use the configured API.
    """
    import re
    headers = _excel_headers(path, sheet)
    numeric = _guess_numeric_headers(path, sheet)
    text = instruction.strip()
    ops: list[dict[str, Any]] = []

    # Row totals: pick likely measure columns while excluding targets/ids/totals.
    if any(k in text for k in ["每行的总数", "每行总数", "逐行合计", "每行合计", "总销售额"]):
        exclude = ("目标", "总", "编号", "序号", "id", "ID", "排名", "完成率")
        sources = [h for h in numeric if not any(x in h for x in exclude)]
        if len(sources) >= 2:
            target_header = next((h for h in headers if "总" in h and h not in sources), "总计")
            ops.append({"op": "row_total", "source": sources, "target": target_header})

    # Group summary: "按区域汇总"
    m = re.search(r"按(.{1,12}?)(?:重新)?汇总", text)
    if m:
        group = m.group(1).strip("的 ")
        group_header = next((h for h in headers if group in h or h in group), "")
        values = [h for h in numeric if "目标" not in h and "编号" not in h and "序号" not in h]
        if group_header and values:
            ops.append({"op": "group_sum", "group_by": group_header, "values": values, "sheet": f"{group_header}汇总"})

    # Extract a named region/category into its own sheet.
    m = re.search(r"(华东|华南|华北|华中|西南|西北|东北).{0,8}(?:单独|筛选|拿出来|提取)", text)
    if m:
        value = m.group(1)
        col = next((h for h in headers if "区域" in h or "地区" in h), "")
        if col:
            ops.append({"op": "filter", "column": col, "equals": value, "sheet": value})

    # Mark under-target rows.
    if any(k in text for k in ["没达标", "未达标", "低于目标", "异常门店", "异常项", "标出来"]):
        total = next((h for h in headers if "总销售" in h or h in {"销售额", "实际", "完成值"}), "")
        target_col = next((h for h in headers if "目标" in h or "预算" in h), "")
        if total and target_col:
            ops.append({"op": "mark_below_target", "value": total, "target": target_col, "status": "达成状态"})

    # Split a workbook by each unique value in a named column.
    if any(k in text for k in ["拆分", "拆开", "拆成", "分成很多", "分成多个"]):
        mentioned = [h for h in headers if h and h in text]
        split_col = max(mentioned, key=len) if mentioned else ""
        if not split_col:
            m = re.search(r"按(?:照)?([^，。\s]{1,12}?)(?:列)?(?:的不同值|不同值|不同元素|内容)?(?:拆分|拆开|拆成)", text)
            if m:
                hint = m.group(1).strip("的 ")
                split_col = next((h for h in headers if hint in h or h in hint), "")
        if split_col:
            mode = "sheets" if any(k in text.lower() for k in ["sheet", "工作表", "同一个excel", "一个excel", "同一文件"]) else "files"
            ops.append({"op": "split_by_column", "column": split_col, "mode": mode})

    # De-duplicate rows, optionally by a named key column.
    if "去重" in text or "重复项" in text or "重复数据" in text:
        subset = [h for h in headers if h and h in text]
        ops.append({"op": "deduplicate", "subset": subset[:3], "keep": "last" if "保留最新" in text or "保留最后" in text else "first"})

    # Fill blanks in a specifically mentioned column.
    if any(k in text for k in ["空白", "空值", "缺失"] ) and any(k in text for k in ["填成", "填为", "补成", "补为"]):
        col = next((h for h in headers if h and h in text), "")
        m = re.search(r"(?:填成|填为|补成|补为)[“\"']?([^”\"'，。]{1,30})", text)
        if col and m:
            ops.append({"op": "fill_missing", "column": col, "value": m.group(1).strip()})

    # Exact value replacement in one named column.
    if "替换" in text:
        col = next((h for h in headers if h and h in text), "")
        m = re.search(r"把?.{0,20}?[“\"']?([^”\"'，。]{1,30})[”\"']?\s*(?:替换成|改成)[“\"']?([^”\"'，。]{1,30})", text)
        if col and m:
            ops.append({"op": "replace", "column": col, "old": m.group(1).strip(), "new": m.group(2).strip()})

    # Common structured derived-column calculations: A/B, A-B, A+B.
    m = re.search(r"([^，。=]{1,20})\s*=\s*([^，。+\-*/]{1,20})\s*([+\-*/])\s*([^，。+\-*/]{1,20})", text)
    if m:
        target_hint, left_hint, symbol, right_hint = [x.strip() for x in m.groups()]
        target_col = next((h for h in headers if target_hint in h or h in target_hint), target_hint)
        left_col = next((h for h in headers if left_hint in h or h in left_hint), "")
        right_col = next((h for h in headers if right_hint in h or h in right_hint), "")
        if left_col and right_col:
            op_map = {"+": "sum", "-": "difference", "*": "product", "/": "ratio"}
            ops.append({"op": "calculate", "source": [left_col, right_col], "target": target_col, "operator": op_map[symbol]})

    if ops or not allow_ai:
        return {"source": "local_rules", "operations": ops, "headers": headers}

    from ..config import load_settings
    cfg = load_settings()
    if not cfg.local_first.ai_for_ambiguous_tasks:
        return {"source": "local_rules", "operations": [], "headers": headers, "note": "本地规则无法确定操作，且已禁用 AI 规划。"}
    try:
        from ..llm import AIClient
        ai = AIClient(timeout_seconds=30, max_retries=0)
        prompt = f"""把用户对 Excel 的要求转换成 JSON 操作计划。不要计算数据，只规划本地程序该做什么。
文件列名：{json.dumps(headers, ensure_ascii=False)}
数值列：{json.dumps(numeric, ensure_ascii=False)}
用户要求：{instruction}

仅允许以下 op：
row_total: {{op, source:[列名], target:列名}}
group_sum: {{op, group_by:列名, values:[数值列], sheet:新表名}}
filter: {{op, column:列名, equals:值, sheet:新表名}}
mark_below_target: {{op, value:实际列, target:目标列, status:状态列}}
sort: {{op, column:列名, ascending:true/false}}
split_by_column: {{op, column:列名, mode:"files"或"sheets"}}
deduplicate: {{op, subset:[列名], keep:"first"或"last"}}
fill_missing: {{op, column:列名, value:填充值}}
replace: {{op, column:列名, old:旧值, new:新值}}
calculate: {{op, source:[列名,列名], target:新列名, operator:"sum"或"difference"或"product"或"ratio"或"percent"}}
modify_where: {{op, target_column:要修改的列, new_value:新值, condition_column:条件列, condition_operator:"equals"或"contains"或"gt"或"gte"或"lt"或"lte"或"is_blank", condition_value:条件值}}

只输出 JSON：{{\"operations\":[...]}}。无法安全映射就返回空 operations。"""
        raw = ai.complete(prompt, system="你是电子表格任务规划器，只输出严格 JSON。", max_output_tokens=1200)
        raw = raw.strip().strip("`")
        if raw.startswith("json"):
            raw = raw[4:].strip()
        parsed = json.loads(raw)
        return {"source": "ai_plan", "operations": parsed.get("operations", []), "headers": headers}
    except Exception as exc:
        return {"source": "local_rules", "operations": [], "headers": headers, "note": f"AI规划不可用：{exc}"}


def excel_apply_plan(path: str, plan_json: str | dict[str, Any], output_path: str = "", sheet: str = "") -> dict[str, Any]:
    """Execute a restricted spreadsheet plan locally with openpyxl/pandas."""
    import shutil
    from openpyxl import load_workbook
    from openpyxl.styles import PatternFill

    src = resolve_user_path(path)
    target = resolve_user_path(output_path) if output_path else src.with_name(src.stem + "_已处理" + src.suffix)
    target.parent.mkdir(parents=True, exist_ok=True)
    if src != target:
        shutil.copy2(src, target)
    plan = json.loads(plan_json) if isinstance(plan_json, str) else plan_json
    operations = list(plan.get("operations", []))
    executed: list[str] = []
    extra_outputs: list[str] = []

    for op in operations:
        kind = str(op.get("op", ""))
        if kind == "row_total":
            result = excel_calculate_row_totals(str(target), ",".join(op.get("source", [])), str(op.get("target", "总计")), sheet)
            executed.append(f"逐行合计 {result['rows_updated']} 行 → {op.get('target','总计')}")
        elif kind == "group_sum":
            df = pd.read_excel(target, sheet_name=sheet if sheet else 0)
            group = str(op["group_by"]); values = [str(v) for v in op.get("values", []) if str(v) in df.columns]
            if group not in df.columns or not values:
                raise ValueError("group_sum 列名不匹配")
            out = df.groupby(group, dropna=False)[values].sum(numeric_only=True).reset_index()
            with pd.ExcelWriter(target, engine="openpyxl", mode="a", if_sheet_exists="replace") as writer:
                out.to_excel(writer, sheet_name=str(op.get("sheet") or f"{group}汇总")[:31], index=False)
            executed.append(f"按 {group} 汇总 → {op.get('sheet') or group+'汇总'}")
        elif kind == "filter":
            df = pd.read_excel(target, sheet_name=sheet if sheet else 0)
            col = str(op["column"]); value = op.get("equals")
            if col not in df.columns:
                raise ValueError(f"filter 列不存在：{col}")
            out = df[df[col].astype(str) == str(value)]
            with pd.ExcelWriter(target, engine="openpyxl", mode="a", if_sheet_exists="replace") as writer:
                out.to_excel(writer, sheet_name=str(op.get("sheet") or str(value))[:31], index=False)
            executed.append(f"筛选 {col}={value} → {op.get('sheet') or value}")
        elif kind == "mark_below_target":
            wb = load_workbook(target)
            ws = wb[sheet] if sheet else wb[wb.sheetnames[0]]
            headers = {str(c.value).strip(): c.column for c in ws[1] if c.value is not None}
            value_col = headers.get(str(op["value"])); target_col = headers.get(str(op["target"]))
            status_name = str(op.get("status") or "达成状态")
            status_col = headers.get(status_name)
            if not status_col:
                status_col = ws.max_column + 1; ws.cell(1, status_col).value = status_name
            if not value_col or not target_col:
                raise ValueError("mark_below_target 列名不匹配")
            red = PatternFill("solid", fgColor="FFF2CC")
            count = 0
            for r in range(2, ws.max_row + 1):
                a = ws.cell(r, value_col).value; b = ws.cell(r, target_col).value
                try:
                    below = float(a) < float(b)
                except (TypeError, ValueError):
                    continue
                ws.cell(r, status_col).value = "未达标" if below else "达标"
                if below:
                    count += 1
                    for c in range(1, ws.max_column + 1):
                        ws.cell(r, c).fill = red
            wb.save(target)
            executed.append(f"标记未达标 {count} 行")
        elif kind == "sort":
            df = pd.read_excel(target, sheet_name=sheet if sheet else 0)
            col = str(op["column"])
            if col not in df.columns:
                raise ValueError(f"sort 列不存在：{col}")
            out = df.sort_values(col, ascending=bool(op.get("ascending", True)))
            target_sheet = sheet or pd.ExcelFile(target).sheet_names[0]
            with pd.ExcelWriter(target, engine="openpyxl", mode="a", if_sheet_exists="replace") as writer:
                out.to_excel(writer, sheet_name=target_sheet[:31], index=False)
            executed.append(f"按 {col} 排序")
        elif kind == "split_by_column":
            result = excel_split_by_column(str(target), str(op["column"]), mode=str(op.get("mode", "files")), sheet=sheet)
            extra_outputs.extend(result.get("created", []))
            executed.append(f"按 {op['column']} 拆分为 {result['groups']} 组")
        elif kind == "deduplicate":
            df = pd.read_excel(target, sheet_name=sheet if sheet else 0)
            subset = [str(c) for c in op.get("subset", []) if str(c) in df.columns] or None
            before = len(df)
            out = df.drop_duplicates(subset=subset, keep=str(op.get("keep", "first")) if str(op.get("keep", "first")) in {"first", "last"} else "first")
            target_sheet = sheet or pd.ExcelFile(target).sheet_names[0]
            with pd.ExcelWriter(target, engine="openpyxl", mode="a", if_sheet_exists="replace") as writer:
                out.to_excel(writer, sheet_name=target_sheet[:31], index=False)
            executed.append(f"去重 {before-len(out)} 行")
        elif kind == "fill_missing":
            df = pd.read_excel(target, sheet_name=sheet if sheet else 0)
            col = str(op["column"])
            if col not in df.columns:
                raise ValueError(f"fill_missing 列不存在：{col}")
            mask = df[col].isna() | (df[col].astype(str).str.strip() == "")
            count = int(mask.sum()); df.loc[mask, col] = op.get("value")
            target_sheet = sheet or pd.ExcelFile(target).sheet_names[0]
            with pd.ExcelWriter(target, engine="openpyxl", mode="a", if_sheet_exists="replace") as writer:
                df.to_excel(writer, sheet_name=target_sheet[:31], index=False)
            executed.append(f"填充 {col} 空值 {count} 行")
        elif kind == "replace":
            df = pd.read_excel(target, sheet_name=sheet if sheet else 0)
            col = str(op["column"])
            if col not in df.columns:
                raise ValueError(f"replace 列不存在：{col}")
            mask = df[col].astype(str) == str(op.get("old")); count = int(mask.sum()); df.loc[mask, col] = op.get("new")
            target_sheet = sheet or pd.ExcelFile(target).sheet_names[0]
            with pd.ExcelWriter(target, engine="openpyxl", mode="a", if_sheet_exists="replace") as writer:
                df.to_excel(writer, sheet_name=target_sheet[:31], index=False)
            executed.append(f"替换 {col} {count} 行")
        elif kind == "calculate":
            df = pd.read_excel(target, sheet_name=sheet if sheet else 0)
            sources = [str(c) for c in op.get("source", [])]
            if any(c not in df.columns for c in sources) or not sources:
                raise ValueError("calculate 源列不匹配")
            nums = [pd.to_numeric(df[c], errors="coerce") for c in sources]
            operator = str(op.get("operator", "sum")); target_col = str(op.get("target", "计算结果"))
            if operator in {"sum", "add"}: result = pd.concat(nums, axis=1).sum(axis=1, min_count=1)
            elif operator in {"product", "multiply"}: result = pd.concat(nums, axis=1).prod(axis=1, min_count=1)
            elif operator in {"difference", "subtract"} and len(nums) == 2: result = nums[0] - nums[1]
            elif operator in {"ratio", "divide"} and len(nums) == 2: result = nums[0] / nums[1].replace(0, pd.NA)
            elif operator in {"percent", "percentage"} and len(nums) == 2: result = nums[0] / nums[1].replace(0, pd.NA) * 100
            else: raise ValueError(f"calculate 运算不支持：{operator}")
            df[target_col] = result.round(6)
            target_sheet = sheet or pd.ExcelFile(target).sheet_names[0]
            with pd.ExcelWriter(target, engine="openpyxl", mode="a", if_sheet_exists="replace") as writer:
                df.to_excel(writer, sheet_name=target_sheet[:31], index=False)
            executed.append(f"计算新列 {target_col}")
        elif kind == "modify_where":
            df = pd.read_excel(target, sheet_name=sheet if sheet else 0)
            target_col = str(op["target_column"]); cond_col = str(op["condition_column"])
            if target_col not in df.columns or cond_col not in df.columns:
                raise ValueError("modify_where 列名不匹配")
            mask = _excel_condition_mask(df[cond_col], str(op.get("condition_operator", "equals")), op.get("condition_value"))
            count = int(mask.sum()); df.loc[mask, target_col] = op.get("new_value")
            target_sheet = sheet or pd.ExcelFile(target).sheet_names[0]
            with pd.ExcelWriter(target, engine="openpyxl", mode="a", if_sheet_exists="replace") as writer:
                df.to_excel(writer, sheet_name=target_sheet[:31], index=False)
            executed.append(f"条件修改 {target_col} {count} 行")
        else:
            raise ValueError(f"不允许的 Excel 操作：{kind}")
    record("excel_apply_plan", {"path": str(src), "output": str(target), "ops": len(executed)}, detail=" | ".join(executed))
    return {"file": str(target), "executed": executed, "extra_outputs": extra_outputs}


def excel_process_instruction(path: str, instruction: str, output_path: str = "", sheet: str = "") -> dict[str, Any]:
    plan = excel_plan_from_instruction(path, instruction, sheet=sheet, allow_ai=True)
    if not plan.get("operations"):
        return {"ok": False, "plan": plan, "message": "没有得到可安全执行的本地 Excel 操作计划。"}
    result = excel_apply_plan(path, plan, output_path=output_path, sheet=sheet)
    result["plan_source"] = plan.get("source", "")
    result["local_analysis"] = excel_local_analysis(result["file"], sheet=sheet)
    return {"ok": True, **result}


def excel_write_analysis_sheet(path: str, analysis: str, sheet_name: str = "AI分析") -> str:
    """Write analysis text into a dedicated worksheet locally."""
    from openpyxl import load_workbook
    from openpyxl.styles import Alignment, Font
    target = resolve_user_path(path)
    wb = load_workbook(target)
    if sheet_name in wb.sheetnames:
        ws = wb[sheet_name]
        ws.delete_rows(1, ws.max_row)
    else:
        ws = wb.create_sheet(sheet_name[:31])
    ws["A1"] = "小智本地分析"
    ws["A1"].font = Font(bold=True, size=14)
    lines = [x.strip() for x in str(analysis).splitlines() if x.strip()]
    for i, line in enumerate(lines, start=3):
        ws.cell(i, 1).value = line
        ws.cell(i, 1).alignment = Alignment(wrap_text=True, vertical="top")
    ws.column_dimensions["A"].width = 100
    wb.save(target)
    record("excel_write_analysis_sheet", {"path": str(target), "sheet": sheet_name}, detail=f"lines={len(lines)}")
    return str(target)


def excel_merge_files(paths: list[str], output_path: str, sheet_name: str = "合并数据") -> str:
    """Append rows from multiple Excel files into one local workbook, adding 来源文件."""
    frames = []
    resolved = []
    for path in paths:
        p = resolve_user_path(path)
        df = pd.read_excel(p, sheet_name=0)
        df.insert(0, "来源文件", p.name)
        frames.append(df)
        resolved.append(str(p))
    if not frames:
        raise ValueError("没有可合并的 Excel 文件")
    merged = pd.concat(frames, ignore_index=True, sort=False)
    target = resolve_user_path(output_path)
    if target.suffix.lower() != ".xlsx":
        target = target.with_suffix(".xlsx")
    target.parent.mkdir(parents=True, exist_ok=True)
    with pd.ExcelWriter(target, engine="openpyxl") as writer:
        merged.to_excel(writer, sheet_name=sheet_name[:31], index=False)
    record("excel_merge_files", {"paths": resolved, "output": str(target)}, detail=f"rows={len(merged)}")
    return str(target)


def _safe_excel_token(value: Any, limit: int = 48) -> str:
    import re
    text = re.sub(r'[\\/:*?"<>|\r\n\[\]]+', '_', str(value)).strip(' ._')
    return (text or '空白')[:limit]


def _style_dataframe_workbook(path: Path) -> None:
    """Apply lightweight readable formatting to generated tabular workbooks."""
    from openpyxl import load_workbook
    from openpyxl.styles import Font, PatternFill, Alignment
    wb = load_workbook(path)
    for ws in wb.worksheets:
        if ws.max_row < 1:
            continue
        for cell in ws[1]:
            cell.font = Font(bold=True)
            cell.fill = PatternFill('solid', fgColor='D9EAF7')
            cell.alignment = Alignment(horizontal='center', vertical='center')
        ws.freeze_panes = 'A2'
        ws.auto_filter.ref = ws.dimensions
        for col_cells in ws.iter_cols(min_row=1, max_row=min(ws.max_row, 200)):
            width = min(max(len(str(c.value or '')) for c in col_cells) + 2, 36)
            ws.column_dimensions[col_cells[0].column_letter].width = max(width, 10)
    wb.save(path)


def excel_split_by_column(
    path: str,
    column: str,
    mode: str = 'files',
    output_dir: str = '',
    output_path: str = '',
    sheet: str = '',
    include_blank: bool = False,
    max_groups: int = 200,
) -> dict[str, Any]:
    """Split one worksheet by each unique value in a column, entirely locally.

    mode='files' creates one XLSX per unique value; mode='sheets' creates one workbook
    containing one worksheet per unique value. A manifest is always generated.
    """
    src = resolve_user_path(path)
    df = pd.read_excel(src, sheet_name=sheet if sheet else 0)
    if column not in df.columns:
        raise ValueError(f'拆分列不存在：{column}')
    work = df.copy()
    if include_blank:
        keys = work[column].where(work[column].notna() & (work[column].astype(str).str.strip() != ''), '空白').astype(str)
    else:
        mask = work[column].notna() & (work[column].astype(str).str.strip() != '')
        work = work.loc[mask].copy()
        keys = work[column].astype(str)
    values = list(dict.fromkeys(str(v) for v in keys.tolist()))
    if len(values) > max(1, int(max_groups)):
        raise ValueError(f'唯一值数量 {len(values)} 超过安全上限 {max_groups}，请先筛选后再拆分。')
    if not values:
        raise ValueError('拆分列没有可用值。')
    base_sheet = sheet or pd.ExcelFile(src).sheet_names[0]
    created: list[str] = []
    manifest_rows: list[dict[str, Any]] = []
    mode = mode.lower().strip()

    if mode in {'sheet', 'sheets', '工作表', '同一文件'}:
        target = resolve_user_path(output_path) if output_path else src.with_name(f'{src.stem}_按{_safe_excel_token(column)}拆分.xlsx')
        if target.suffix.lower() != '.xlsx':
            target = target.with_suffix('.xlsx')
        target.parent.mkdir(parents=True, exist_ok=True)
        used: set[str] = set()
        with pd.ExcelWriter(target, engine='openpyxl') as writer:
            for value in values:
                out = work.loc[keys == value]
                name = _safe_excel_token(value, 28) or '空白'
                candidate = name
                n = 2
                while candidate in used:
                    candidate = f'{name[:25]}_{n}'
                    n += 1
                used.add(candidate)
                out.to_excel(writer, sheet_name=candidate[:31], index=False)
                manifest_rows.append({'拆分值': value, '行数': len(out), '工作表': candidate[:31]})
            pd.DataFrame(manifest_rows).to_excel(writer, sheet_name='拆分清单', index=False)
        _style_dataframe_workbook(target)
        created.append(str(target))
    elif mode in {'file', 'files', '文件'}:
        out_dir = resolve_user_path(output_dir) if output_dir else src.parent / f'{src.stem}_按{_safe_excel_token(column)}拆分'
        out_dir.mkdir(parents=True, exist_ok=True)
        used_names: set[str] = set()
        for value in values:
            out = work.loc[keys == value]
            stem = _safe_excel_token(value)
            candidate = stem
            n = 2
            while candidate.lower() in used_names:
                candidate = f'{stem}_{n}'
                n += 1
            used_names.add(candidate.lower())
            target = out_dir / f'{candidate}.xlsx'
            with pd.ExcelWriter(target, engine='openpyxl') as writer:
                out.to_excel(writer, sheet_name=_safe_excel_token(base_sheet, 31), index=False)
            _style_dataframe_workbook(target)
            created.append(str(target))
            manifest_rows.append({'拆分值': value, '行数': len(out), '文件': target.name})
        manifest = out_dir / '拆分清单.xlsx'
        pd.DataFrame(manifest_rows).to_excel(manifest, index=False)
        _style_dataframe_workbook(manifest)
        created.append(str(manifest))
    else:
        raise ValueError('mode 只支持 files 或 sheets。')
    record('excel_split_by_column', {'path': str(src), 'column': column, 'mode': mode, 'groups': len(values)}, detail=' | '.join(created[:20]))
    return {'source': str(src), 'column': column, 'groups': len(values), 'created': created, 'manifest': manifest_rows}


def _excel_condition_mask(series: pd.Series, operator: str, value: Any) -> pd.Series:
    op = operator.lower().strip()
    if op == 'equals':
        return series.astype(str) == str(value)
    if op == 'not_equals':
        return series.astype(str) != str(value)
    if op == 'contains':
        return series.astype(str).str.contains(str(value), case=False, na=False, regex=False)
    if op == 'not_contains':
        return ~series.astype(str).str.contains(str(value), case=False, na=False, regex=False)
    if op == 'is_blank':
        return series.isna() | (series.astype(str).str.strip() == '')
    if op == 'not_blank':
        return series.notna() & (series.astype(str).str.strip() != '')
    left = pd.to_numeric(series, errors='coerce')
    right = float(value)
    if op == 'gt':
        return left > right
    if op == 'gte':
        return left >= right
    if op == 'lt':
        return left < right
    if op == 'lte':
        return left <= right
    raise ValueError(f'不支持的条件运算：{operator}')


def excel_modify_rows(
    path: str,
    target_column: str,
    new_value: Any,
    condition_column: str = '',
    condition_operator: str = 'equals',
    condition_value: Any = '',
    output_path: str = '',
    sheet: str = '',
) -> dict[str, Any]:
    """Modify one column only for rows matching a structured condition."""
    src = resolve_user_path(path)
    df = pd.read_excel(src, sheet_name=sheet if sheet else 0)
    if target_column not in df.columns:
        raise ValueError(f'目标列不存在：{target_column}')
    if condition_column:
        if condition_column not in df.columns:
            raise ValueError(f'条件列不存在：{condition_column}')
        mask = _excel_condition_mask(df[condition_column], condition_operator, condition_value)
    else:
        mask = pd.Series(True, index=df.index)
    count = int(mask.sum())
    df.loc[mask, target_column] = new_value
    target = resolve_user_path(output_path) if output_path else src.with_name(src.stem + '_已修改.xlsx')
    target.parent.mkdir(parents=True, exist_ok=True)
    with pd.ExcelWriter(target, engine='openpyxl') as writer:
        df.to_excel(writer, sheet_name=(sheet or '数据')[:31], index=False)
    _style_dataframe_workbook(target)
    record('excel_modify_rows', {'path': str(src), 'output': str(target), 'rows': count, 'target_column': target_column})
    return {'file': str(target), 'rows_modified': count}


def excel_replace_values(path: str, column: str, old_value: Any, new_value: Any, output_path: str = '', sheet: str = '') -> dict[str, Any]:
    src = resolve_user_path(path)
    df = pd.read_excel(src, sheet_name=sheet if sheet else 0)
    if column not in df.columns:
        raise ValueError(f'列不存在：{column}')
    mask = df[column].astype(str) == str(old_value)
    count = int(mask.sum())
    df.loc[mask, column] = new_value
    target = resolve_user_path(output_path) if output_path else src.with_name(src.stem + '_已替换.xlsx')
    with pd.ExcelWriter(target, engine='openpyxl') as writer:
        df.to_excel(writer, sheet_name=(sheet or '数据')[:31], index=False)
    _style_dataframe_workbook(target)
    return {'file': str(target), 'replaced': count}


def excel_fill_missing(path: str, column: str, value: Any, output_path: str = '', sheet: str = '') -> dict[str, Any]:
    src = resolve_user_path(path)
    df = pd.read_excel(src, sheet_name=sheet if sheet else 0)
    if column not in df.columns:
        raise ValueError(f'列不存在：{column}')
    mask = df[column].isna() | (df[column].astype(str).str.strip() == '')
    count = int(mask.sum())
    df.loc[mask, column] = value
    target = resolve_user_path(output_path) if output_path else src.with_name(src.stem + '_已填空.xlsx')
    with pd.ExcelWriter(target, engine='openpyxl') as writer:
        df.to_excel(writer, sheet_name=(sheet or '数据')[:31], index=False)
    _style_dataframe_workbook(target)
    return {'file': str(target), 'filled': count}


def excel_deduplicate(path: str, subset: str = '', keep: str = 'first', output_path: str = '', sheet: str = '') -> dict[str, Any]:
    src = resolve_user_path(path)
    df = pd.read_excel(src, sheet_name=sheet if sheet else 0)
    cols = [x.strip() for x in subset.replace('，', ',').split(',') if x.strip()] if subset else None
    if cols:
        missing = [c for c in cols if c not in df.columns]
        if missing:
            raise ValueError(f'去重列不存在：{missing}')
    before = len(df)
    keep_value: str | bool = keep if keep in {'first', 'last'} else 'first'
    out = df.drop_duplicates(subset=cols, keep=keep_value)
    target = resolve_user_path(output_path) if output_path else src.with_name(src.stem + '_已去重.xlsx')
    with pd.ExcelWriter(target, engine='openpyxl') as writer:
        out.to_excel(writer, sheet_name=(sheet or '数据')[:31], index=False)
    _style_dataframe_workbook(target)
    return {'file': str(target), 'removed': before - len(out), 'rows': len(out)}


def excel_calculate_column(
    path: str,
    source_columns: str,
    target_column: str,
    operator: str = 'sum',
    output_path: str = '',
    sheet: str = '',
) -> dict[str, Any]:
    """Calculate a derived column using a fixed safe operator, never arbitrary eval/code."""
    src = resolve_user_path(path)
    df = pd.read_excel(src, sheet_name=sheet if sheet else 0)
    cols = [x.strip() for x in source_columns.replace('，', ',').split(',') if x.strip()]
    missing = [c for c in cols if c not in df.columns]
    if missing or not cols:
        raise ValueError(f'计算列不存在或为空：{missing or cols}')
    nums = [pd.to_numeric(df[c], errors='coerce') for c in cols]
    op = operator.lower().strip()
    if op in {'sum', 'add'}:
        result = pd.concat(nums, axis=1).sum(axis=1, min_count=1)
    elif op in {'product', 'multiply'}:
        result = pd.concat(nums, axis=1).prod(axis=1, min_count=1)
    elif op in {'difference', 'subtract'}:
        if len(nums) != 2:
            raise ValueError('difference 需要两列。')
        result = nums[0] - nums[1]
    elif op in {'ratio', 'divide'}:
        if len(nums) != 2:
            raise ValueError('ratio 需要两列。')
        result = nums[0] / nums[1].replace(0, pd.NA)
    elif op in {'percent', 'percentage'}:
        if len(nums) != 2:
            raise ValueError('percent 需要两列。')
        result = nums[0] / nums[1].replace(0, pd.NA) * 100
    else:
        raise ValueError(f'不支持的计算操作：{operator}')
    df[target_column] = result.round(6)
    target = resolve_user_path(output_path) if output_path else src.with_name(src.stem + '_已计算.xlsx')
    with pd.ExcelWriter(target, engine='openpyxl') as writer:
        df.to_excel(writer, sheet_name=(sheet or '数据')[:31], index=False)
    _style_dataframe_workbook(target)
    return {'file': str(target), 'target_column': target_column, 'rows': len(df)}


def excel_lookup_merge(
    left_path: str,
    right_path: str,
    left_key: str,
    right_key: str = '',
    right_columns: str = '',
    output_path: str = '',
    how: str = 'left',
) -> dict[str, Any]:
    """VLOOKUP-style local join between two workbooks using pandas merge."""
    left = resolve_user_path(left_path)
    right = resolve_user_path(right_path)
    ldf = pd.read_excel(left, sheet_name=0)
    rdf = pd.read_excel(right, sheet_name=0)
    rkey = right_key or left_key
    if left_key not in ldf.columns or rkey not in rdf.columns:
        raise ValueError('匹配键列不存在。')
    cols = [x.strip() for x in right_columns.replace('，', ',').split(',') if x.strip()]
    if cols:
        missing = [c for c in cols if c not in rdf.columns]
        if missing:
            raise ValueError(f'右表字段不存在：{missing}')
        rdf = rdf[[rkey] + [c for c in cols if c != rkey]]
    before = len(ldf)
    join_how = how if how in {'left', 'inner', 'right', 'outer'} else 'left'
    out = ldf.merge(rdf, how=join_how, left_on=left_key, right_on=rkey, suffixes=('', '_匹配'))
    if rkey != left_key and rkey in out.columns:
        out.drop(columns=[rkey], inplace=True)
    target = resolve_user_path(output_path) if output_path else left.with_name(left.stem + '_匹配结果.xlsx')
    with pd.ExcelWriter(target, engine='openpyxl') as writer:
        out.to_excel(writer, sheet_name='匹配结果', index=False)
    _style_dataframe_workbook(target)
    record('excel_lookup_merge', {'left': str(left), 'right': str(right), 'output': str(target), 'rows': len(out)})
    return {'file': str(target), 'left_rows': before, 'result_rows': len(out)}


def excel_merge_folder(input_dir: str, output_path: str = '', recursive: bool = False, add_source_column: bool = True, max_files: int = 200) -> dict[str, Any]:
    folder = resolve_user_path(input_dir)
    pattern = '**/*.xlsx' if recursive else '*.xlsx'
    files = [p for p in folder.glob(pattern) if p.is_file() and not p.name.startswith('~$')][:max_files]
    if not files:
        raise ValueError('目录中没有找到 .xlsx 文件。')
    frames = []
    for p in files:
        df = pd.read_excel(p, sheet_name=0)
        if add_source_column:
            df.insert(0, '来源文件', p.name)
        frames.append(df)
    out = pd.concat(frames, ignore_index=True, sort=False)
    target = resolve_user_path(output_path) if output_path else folder / '合并总表.xlsx'
    with pd.ExcelWriter(target, engine='openpyxl') as writer:
        out.to_excel(writer, sheet_name='合并数据', index=False)
    _style_dataframe_workbook(target)
    return {'file': str(target), 'files_merged': len(files), 'rows': len(out)}


def excel_batch_process_directory(
    input_dir: str,
    instruction: str,
    output_dir: str = '',
    recursive: bool = False,
    max_files: int = 100,
) -> dict[str, Any]:
    """Apply one spreadsheet instruction repeatedly to a folder of XLSX files.

    The plan is created once from the first workbook and reused for every file, so an
    ambiguous instruction never causes one AI call per workbook.
    """
    folder = resolve_user_path(input_dir)
    pattern = '**/*.xlsx' if recursive else '*.xlsx'
    files = [p for p in folder.glob(pattern) if p.is_file() and not p.name.startswith('~$')][:max_files]
    if not files:
        raise ValueError('目录中没有找到 .xlsx 文件。')
    plan = excel_plan_from_instruction(str(files[0]), instruction, allow_ai=True)
    if not plan.get('operations'):
        return {'ok': False, 'plan': plan, 'message': '没有得到可安全复用的批处理计划。'}
    out_dir = resolve_user_path(output_dir) if output_dir else folder / '小智批量处理结果'
    out_dir.mkdir(parents=True, exist_ok=True)
    ok: list[dict[str, Any]] = []
    failed: list[dict[str, str]] = []
    for src in files:
        try:
            target = out_dir / src.name
            result = excel_apply_plan(str(src), plan, output_path=str(target))
            ok.append({'source': str(src), 'output': result.get('file'), 'extra_outputs': result.get('extra_outputs', [])})
        except Exception as exc:
            failed.append({'source': str(src), 'error': str(exc)})
    record('excel_batch_process_directory', {'input_dir': str(folder), 'files': len(files), 'ok': len(ok), 'failed': len(failed)}, detail=str(plan)[:1500])
    return {'ok': True, 'plan_source': plan.get('source'), 'processed': len(ok), 'failed_count': len(failed), 'output_dir': str(out_dir), 'results': ok, 'failed': failed}
