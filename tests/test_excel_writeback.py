from pathlib import Path
from openpyxl import Workbook, load_workbook
from xiaozhi_assistant.tools.office import excel_calculate_row_totals, excel_write_analysis


def test_excel_writeback(tmp_path: Path):
    p = tmp_path / "demo.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.append(["门店","服饰","数码","家居","食品","美妆","总销售额（万元）","月度目标（万元）"])
    ws.append(["A",1,2,3,4,5,None,20])
    ws.append(["B",2,2,2,2,2,None,9])
    ws.append([])
    ws.append(["AI 分析区"])
    ws.append(["销售前三",None,None])
    ws.append(["低于目标门店",None,None])
    ws.append(["品类表现",None,None])
    ws.append(["一句话总结",None,None])
    wb.save(p)
    result = excel_calculate_row_totals(str(p))
    assert result["rows_updated"] == 2
    excel_write_analysis(str(p), "A", "B", "服饰", "总结")
    wb2 = load_workbook(p)
    ws2 = wb2.active
    assert ws2["G2"].value == 15
    assert ws2["G3"].value == 10
    assert ws2["C6"].value == "A"
    assert ws2["C9"].value == "总结"
