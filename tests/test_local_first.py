from pathlib import Path

from openpyxl import Workbook, load_workbook

from xiaozhi_assistant.local_text import extractive_summary, looks_like_work_task
from xiaozhi_assistant.tools.office import excel_plan_from_instruction, excel_process_instruction
from xiaozhi_assistant.tools.workmate import _parse_clock


def test_local_summary_and_task_detection():
    text = "第一季度销售增长明显。华东区域表现最好。西南区域仍未完成目标。建议下周复盘门店转化。"
    summary = extractive_summary(text, max_sentences=2)
    assert summary
    assert looks_like_work_task("把附件销售表按区域汇总一下，下午发我")


def test_excel_process_uses_local_rules(tmp_path: Path):
    p = tmp_path / "sales.xlsx"
    wb = Workbook(); ws = wb.active
    ws.append(["门店", "区域", "服饰", "数码", "月度目标（万元）", "总销售额（万元）"])
    ws.append(["A", "华东", 10, 20, 40, None])
    ws.append(["B", "华南", 8, 7, 20, None])
    wb.save(p)
    plan = excel_plan_from_instruction(str(p), "把每行总销售额算出来", allow_ai=False)
    assert plan["source"] == "local_rules"
    assert plan["operations"][0]["op"] == "row_total"
    result = excel_process_instruction(str(p), "把每行总销售额算出来")
    assert result["ok"] is True
    assert result["plan_source"] == "local_rules"
    out = load_workbook(result["file"]).active
    assert out["F2"].value == 30
    assert out["F3"].value == 15


def test_parse_chinese_schedule_time():
    from datetime import datetime
    now = datetime.fromisoformat("2026-09-08T09:00:00+08:00")
    dt = _parse_clock("今天下午两点半发给老板", now=now)
    assert dt is not None
    assert dt.hour == 14 and dt.minute == 30
